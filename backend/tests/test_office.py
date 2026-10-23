import asyncio
import io
import json
import time
import zipfile
import pytest
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
from app import office, office_formats as fmt, file_tools, provider, data_transfer
from app.db import Conversation, Run, Session
from .test_app import API
from .test_app import wait_run


def wait_file(client, key):
    for _ in range(150):
        row = client.get(f"{API}/files/{key}").json()
        if row["file_status"] not in {"queued", "parsing"}:
            return row
        time.sleep(0.05)
    pytest.fail("File parser did not finish")


def test_all_output_formats_open_and_formula_caches_recalculate():
    blocks = [
        {"kind": "heading", "text": "项目进展"},
        {"kind": "paragraph", "text": "中文办公报告：计划与结果。"},
        {"kind": "table", "rows": [["事项", "数量"], ["完成", "3"]]},
    ]
    doc = fmt.document("docx", "项目报告", blocks, "内部资料", "XiaoXi")
    parsed = fmt.parse(doc, "报告.docx")
    assert any("中文办公报告" in c["text"] for c in parsed["chunks"])
    assert len(Document(io.BytesIO(doc)).tables) == 1
    pdf = fmt.document("pdf", "项目报告", blocks)
    assert "中文办公报告" in PdfReader(io.BytesIO(pdf)).pages[0].extract_text()
    assert fmt.pdf_page(pdf, 1)["image"].startswith("data:image/png;base64,")
    for a, b, expected in [(2, 3, 5), (7, 4, 11)]:
        raw = fmt.spreadsheet(
            [
                {
                    "name": "明细",
                    "rows": [
                        ["项目", "金额"],
                        ["甲", a],
                        ["乙", b],
                        ["合计", {"formula": "SUM(B2:B3)", "number_format": "0.00"}],
                        ["校验", {"formula": "B4*2"}],
                    ],
                }
            ]
        )
        wb = load_workbook(io.BytesIO(raw), data_only=True)
        assert wb["明细"]["B4"].value == expected
        assert wb["明细"]["B5"].value == expected * 2
        wb.close()
        wb = load_workbook(io.BytesIO(raw), data_only=False)
        assert wb["明细"]["B4"].value == "=SUM(B2:B3)"
        wb.close()
    for kind in ["txt", "md", "csv", "json"]:
        source = (
            '{"项目":"测试"}'
            if kind == "json"
            else '项目,内容\n测试,"多行\n中文"'
            if kind == "csv"
            else "# 中文资料\n正文"
        )
        assert fmt.parse(fmt.text_file(kind, source), "文件." + kind)["chunks"]


def test_parser_limits_and_format_failures():
    with pytest.raises(ValueError, match="加密"):
        from pypdf import PdfWriter

        writer = PdfWriter()
        writer.add_blank_page(200, 200)
        writer.encrypt("secret")
        data = io.BytesIO()
        writer.write(data)
        fmt.parse(data.getvalue(), "加密.pdf")
    from reportlab.pdfgen import canvas

    blank = io.BytesIO()
    c = canvas.Canvas(blank)
    c.showPage()
    c.save()
    with pytest.raises(ValueError, match="OCR"):
        fmt.parse(blank.getvalue(), "扫描.pdf")
    with pytest.raises(ValueError):
        fmt.parse(b"fake", "fake.pdf")
    with pytest.raises(ValueError):
        fmt.parse(b"abc", "test.exe")
    with pytest.raises(ValueError):
        fmt.text_file("json", "{bad")
    assert "'=HYPERLINK" in fmt.text_file("csv", '=HYPERLINK("x")').decode("utf-8-sig")
    for formula in [
        "A1",
        'WEBSERVICE("https://example.com")',
        "1/0",
        "SUM([external.xlsx]Sheet1!A1)",
    ]:
        with pytest.raises((ValueError, ZeroDivisionError)):
            fmt.spreadsheet([{"name": "Sheet1", "rows": [[{"formula": formula}]]}])
    payload = io.BytesIO()
    with zipfile.ZipFile(payload, "w") as z:
        z.writestr("word/vbaProject.bin", b"macro")
    with pytest.raises(ValueError, match="宏"):
        fmt.parse(payload.getvalue(), "macro.docx")


def test_upload_parse_claim_tool_generation_preview_and_portable_data(client):
    uploaded = client.post(
        API + "/files",
        files={
            "file": (
                "资料.docx",
                fmt.document(
                    "docx",
                    "项目资料",
                    [{"kind": "paragraph", "text": "预算100元，完成3项任务。"}],
                ),
                "application/octet-stream",
            )
        },
    )
    assert uploaded.status_code == 200
    key = uploaded.json()["id"]
    assert wait_file(client, key)["file_status"] in {"ready", "partial"}
    with Session() as db:
        conversation = Conversation(title="办公验收")
        db.add(conversation)
        db.flush()
        office.claim(db, [key], conversation.id)
        run = Run(kind="chat", conversation_id=conversation.id, status="running")
        db.add(run)
        db.commit()
        cid, rid = conversation.id, run.id
    assert len(office.list_files(cid)) == 1
    context = provider.user_content("整理资料", [key], {"vision_enabled": False})
    assert "read_file" in context and key in context
    assert "预算100元" not in context
    read = asyncio.run(office.read_file(key, cid))
    assert "预算100元" in json.dumps(read, ensure_ascii=False)
    with pytest.raises(ValueError, match="当前对话"):
        asyncio.run(office.read_file(key, "another"))
    toolmap = {t.name: t for t in file_tools.build(rid, cid)}
    result = json.loads(
        asyncio.run(
            toolmap["create_document"].ainvoke(
                {
                    "name": "报告.docx",
                    "format": "docx",
                    "title": "报告",
                    "blocks": [
                        {"kind": "paragraph", "text": "预算100元，完成3项任务。"}
                    ],
                }
            )
        )
    )
    assert "error" not in result, result
    artifact = result["artifacts"][0]
    response = client.get(f"{API}/artifacts/{artifact['id']}/preview")
    assert response.status_code == 200 and response.json()["parts"]
    assert client.get(f"{API}/artifacts/{artifact['id']}/content").content.startswith(
        b"PK"
    )
    again = json.loads(
        asyncio.run(
            toolmap["create_document"].ainvoke(
                {
                    "name": "报告.docx",
                    "format": "docx",
                    "title": "报告",
                    "blocks": [
                        {"kind": "paragraph", "text": "预算100元，完成3项任务。"}
                    ],
                }
            )
        )
    )
    assert again["artifacts"][0]["id"] == artifact["id"]
    with pytest.raises(ValueError):
        asyncio.run(office.generate(rid, cid, "text", "../escape", "txt", content="x"))
    exported = data_transfer.create_archive()
    with zipfile.ZipFile(exported["path"]) as z:
        assert f"attachments/files/{key}.bin" in z.namelist()
    assert client.get(f"{API}/files/{key}/content").content.startswith(b"PK")
    with Session() as db:
        db.get(Run, rid).status = "completed"
        db.commit()
    assert client.delete(f"{API}/conversations/{cid}").status_code in {200, 204}
    assert client.get(f"{API}/files/{key}").status_code == 404
    assert not (office.root() / f"{key}.bin").exists()
    assert data_transfer.create_archive()["path"]


def test_failed_upload_exposes_error_and_preserves_original(client):
    row = client.post(
        API + "/files", files={"file": ("bad.pdf", b"broken", "application/pdf")}
    ).json()
    result = wait_file(client, row["id"])
    assert result["file_status"] == "failed" and result["error"]
    assert client.get(f"{API}/files/{row['id']}/content").content == b"broken"


def test_agent_uses_uploaded_word_and_pdf_and_publishes_multiple_files(configured):
    client = configured
    ids = []
    for format in ["docx", "pdf"]:
        raw = fmt.document(
            format, "资料", [{"kind": "paragraph", "text": "预算100元，完成3项任务。"}]
        )
        upload = client.post(
            API + "/files",
            files={"file": ("资料." + format, raw, "application/octet-stream")},
        ).json()
        assert wait_file(client, upload["id"])["file_status"] in {"ready", "partial"}
        ids.append(upload["id"])
    conversation = client.post(
        API + "/conversations", json={"title": "office-contract"}
    ).json()
    sent = client.post(
        f"{API}/conversations/{conversation['id']}/messages",
        json={
            "content": "office-contract 请读取资料，生成Word报告和Excel明细表",
            "attachment_ids": ids,
        },
    )
    assert sent.status_code == 200, sent.text
    run = wait_run(client, sent)
    assert run["status"] == "completed", run
    presentation = client.get(f"{API}/runs/{run['id']}/presentation").json()
    events = client.get(f"{API}/runs/{run['id']}/tools").json()
    assert not [
        e["data"].get("result")
        for e in events
        if e.get("data", {}).get("status") == "error"
    ], json.dumps(events, ensure_ascii=True)
    assert {a["name"] for a in presentation["artifacts"]} == {
        "办公报告.docx",
        "项目明细.xlsx",
    }, client.get(f"{API}/runs/{run['id']}/tools").json()
    for artifact in presentation["artifacts"]:
        data = client.get(f"{API}/artifacts/{artifact['id']}/content").content
        if artifact["name"].endswith(".docx"):
            assert "预算100元" in "\n".join(
                p.text for p in Document(io.BytesIO(data)).paragraphs
            )
        else:
            wb = load_workbook(io.BytesIO(data), data_only=True)
            assert wb["明细"]["B4"].value == 100
            wb.close()
