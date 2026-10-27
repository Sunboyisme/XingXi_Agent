"""Pure bounded format adapters; no model, network, shell or user paths."""

import ast
import base64
import csv
import io
import json
import math
import re
import zipfile
from html import escape
from pathlib import Path

LIMIT = 20 * 1024 * 1024
MAX_TEXT = 400_000
MIMES = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pdf": "application/pdf",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "txt": "text/plain",
    "md": "text/markdown",
    "csv": "text/csv",
    "json": "application/json",
}


def check_zip(raw):
    with zipfile.ZipFile(io.BytesIO(raw)) as z:
        entries = z.infolist()
        if len(entries) > 2000 or sum(e.file_size for e in entries) > 40 * 1024 * 1024:
            raise ValueError("文档解压后过大，最多 40 MB / 2000 项")
        if any("vbaproject" in e.filename.lower() for e in entries):
            raise ValueError("不支持带宏的文件，请另存为不含宏的文档")


def parse(raw, name, cached=False):
    extension = Path(name).suffix.lower()[1:]
    if extension == "doc":
        raise ValueError("旧版Word .doc不支持，请在Word中另存为.docx后重新上传。")
    if extension not in MIMES or len(raw) > LIMIT:
        raise ValueError("支持 DOCX、PDF、XLSX、TXT、MD、CSV、JSON，单文件最多 20 MB")
    chunks, warnings = [], []
    total = 0

    def add(text, source, cells=None):
        nonlocal total
        text = str(text)
        if total + len(text) > MAX_TEXT:
            text = text[: max(0, MAX_TEXT - total)]
            if "正文超过 40 万字，已截断" not in warnings:
                warnings.append("正文超过 40 万字，已截断")
        for offset in range(0, len(text), 6000):
            chunks.append(
                {
                    "source": source,
                    "text": text[offset : offset + 6000],
                    **(
                        {"cells": cells}
                        if cells is not None and len(text) <= 6000
                        else {}
                    ),
                }
            )
        total += len(text)

    if extension == "pdf":
        from pypdf import PdfReader

        if not raw.lstrip().startswith(b"%PDF-"):
            raise ValueError("文件内容不是有效 PDF")
        reader = PdfReader(io.BytesIO(raw))
        if reader.is_encrypted:
            raise ValueError("PDF 已加密，请先解密后上传")
        if len(reader.pages) > 150:
            warnings.append("超过 150 页，仅解析前 150 页")
        for index, page in enumerate(reader.pages[:150]):
            if total >= MAX_TEXT:
                break
            text = page.extract_text() or ""
            if text.strip():
                add(text, f"第 {index + 1} 页")
            else:
                warnings.append(f"第 {index + 1} 页没有可提取文字，可能需要 OCR")
    elif extension == "docx":
        from docx import Document
        from docx.table import Table

        check_zip(raw)
        doc = Document(io.BytesIO(raw))
        for index, block in enumerate(doc.iter_inner_content()):
            if total >= MAX_TEXT:
                break
            text = (
                "\n".join(" | ".join(c.text for c in row.cells) for row in block.rows)
                if isinstance(block, Table)
                else block.text
            )
            add(
                text,
                f"正文第 {index + 1} 块"
                + ("（表格）" if isinstance(block, Table) else ""),
            )
        warnings.append("提取正文和表格；复杂版式、图片文字、批注和修订可能不完整")
    elif extension == "xlsx":
        from openpyxl import load_workbook

        check_zip(raw)
        wb = load_workbook(
            io.BytesIO(raw), read_only=True, data_only=cached, keep_links=False
        )
        for sheet in wb.worksheets[:20]:
            if sheet.max_row > 2000 or sheet.max_column > 50:
                warnings.append(f"{sheet.title} 仅提取前 2000 行、50 列")
            for index, row in enumerate(
                sheet.iter_rows(
                    max_row=min(sheet.max_row, 2000),
                    max_col=min(sheet.max_column, 50),
                    values_only=True,
                )
            ):
                if total >= MAX_TEXT:
                    break
                cells = ["" if v is None else str(v) for v in row]
                add(
                    " | ".join(cells),
                    f"工作表 {sheet.title} 第 {index + 1} 行",
                    cells=cells,
                )
        wb.close()
        if not cached:
            warnings.append("输入表格公式仅作为文本读取，不执行公式或外部链接")
    else:
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                text = raw.decode("gb18030")
                warnings.append("按 GB18030 中文编码读取")
            except UnicodeDecodeError:
                raise ValueError("无法识别文字编码，请另存为 UTF-8") from None
        if "\x00" in text:
            raise ValueError("文件包含二进制内容，不是可读取的文本")
        if extension == "json":
            text = json.dumps(
                json.loads(text), ensure_ascii=False, indent=2, allow_nan=False
            )
        if extension == "csv":
            rows = list(csv.reader(io.StringIO(text)))
            for index, row in enumerate(rows[:4000]):
                add(" | ".join(row), f"CSV 第 {index + 1} 行", cells=row)
            if len(rows) > 4000:
                warnings.append("CSV 超过 4000 行，已截断")
        else:
            for index, line in enumerate(text.splitlines()):
                if total >= MAX_TEXT:
                    warnings.append("正文超过 40 万字，已截断")
                    break
                add(line, f"第 {index + 1} 行")
    if not chunks:
        raise ValueError(
            "未提取到可读取文字；扫描文档需要 OCR，空文档或图片不会被当作正文"
        )
    return {
        "chunks": chunks,
        "warnings": list(dict.fromkeys(warnings)),
        "status": "partial" if warnings else "ready",
    }


def validate_blocks(blocks):
    if not isinstance(blocks, list) or not 1 <= len(blocks) <= 400:
        raise ValueError("文档需有 1–400 个内容块")
    if len(json.dumps(blocks, ensure_ascii=False)) > 800_000:
        raise ValueError("文档内容过大")
    for block in blocks:
        if block.get("kind") not in {
            "paragraph",
            "heading",
            "bullet",
            "number",
            "table",
            "page_break",
            "image",
        }:
            raise ValueError("不支持的内容块类型")
        if len(str(block.get("text", ""))) > 12000:
            raise ValueError("单段文字不能超过 12000 字")
        if block["kind"] == "table":
            rows = block.get("rows", [])
            if (
                not rows
                or len(rows) > 200
                or any(not isinstance(r, list) or not 1 <= len(r) <= 12 for r in rows)
            ):
                raise ValueError("表格最多 200 行、12 列")
    return blocks


def document(format, title, blocks, header="", footer=""):
    blocks = validate_blocks(blocks)
    output = io.BytesIO()
    if format == "docx":
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        doc = Document()
        section = doc.sections[0]
        section.top_margin = section.bottom_margin = Inches(0.75)
        section.left_margin = section.right_margin = Inches(0.8)
        for style_name in ["Normal", "Title", "Heading 1", "Heading 2", "Heading 3"]:
            style = doc.styles[style_name]
            style.font.name = "Microsoft YaHei"
            style.element.get_or_add_rPr().rFonts.set(
                qn("w:eastAsia"), "Microsoft YaHei"
            )
        doc.styles["Normal"].font.size = Pt(11)
        doc.styles["Normal"].paragraph_format.space_after = Pt(7)
        for heading in ["Heading 1", "Heading 2", "Heading 3"]:
            doc.styles[heading].font.color.rgb = RGBColor.from_string("303C50")
        if title:
            doc.add_heading(title, 0)
        section.header.paragraphs[0].text = header[:200]
        section.footer.paragraphs[0].text = footer[:200] + "  "
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        section.footer.paragraphs[0]._p.append(field)
        for b in blocks:
            kind, text = b["kind"], str(b.get("text", ""))
            if kind == "heading":
                doc.add_heading(text, max(1, min(3, int(b.get("level", 1)))))
            elif kind == "table":
                rows = b["rows"]
                table = doc.add_table(rows=0, cols=max(map(len, rows)))
                table.style = "Table Grid"
                for index, row in enumerate(rows):
                    cells = table.add_row().cells
                    for cell, value in zip(cells, row):
                        cell.text = str(value)
                        if index == 0:
                            for run in cell.paragraphs[0].runs:
                                run.bold = True
            elif kind == "page_break":
                doc.add_page_break()
            elif kind == "image":
                doc.add_picture(
                    io.BytesIO(base64.b64decode(b["image_data"])),
                    width=Inches(min(6, max(0.5, float(b.get("width_inches", 4))))),
                )
            else:
                doc.add_paragraph(
                    text,
                    style={"bullet": "List Bullet", "number": "List Number"}.get(kind),
                )
        doc.save(output)
    elif format == "pdf":
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib import colors
        from reportlab.platypus import (
            SimpleDocTemplate,
            Paragraph,
            Spacer,
            Table,
            TableStyle,
            PageBreak,
            Image,
        )

        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
        normal = ParagraphStyle(
            "Body",
            fontName="STSong-Light",
            fontSize=11,
            leading=18,
            wordWrap="CJK",
            spaceAfter=9,
        )
        story = []

        def para(text, size=11):
            return Paragraph(
                escape(str(text)).replace("\n", "<br/>"),
                ParagraphStyle(
                    "Text", parent=normal, fontSize=size, leading=size * 1.6
                ),
            )

        if title:
            story.extend([para(title, 21), Spacer(1, 12)])
        list_number = 0
        for b in blocks:
            kind, text = b["kind"], str(b.get("text", ""))
            list_number = list_number + 1 if kind == "number" else 0
            if kind == "table":
                rows = b["rows"]
                count = max(map(len, rows))
                table = Table(
                    [
                        [para(c, 9) for c in row] + [""] * (count - len(row))
                        for row in rows
                    ],
                    colWidths=[480 / count] * count,
                    repeatRows=1,
                    hAlign="LEFT",
                )
                table.setStyle(
                    TableStyle(
                        [
                            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd2db")),
                            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e9eef4")),
                            ("VALIGN", (0, 0), (-1, -1), "TOP"),
                            ("TOPPADDING", (0, 0), (-1, -1), 6),
                            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                        ]
                    )
                )
                story.extend([table, Spacer(1, 12)])
            elif kind == "page_break":
                story.append(PageBreak())
            elif kind == "image":
                from PIL import Image as PillowImage

                raw = base64.b64decode(b["image_data"])
                with PillowImage.open(io.BytesIO(raw)) as picture:
                    width = min(430, float(b.get("width_inches", 4)) * 72)
                    height = width * picture.height / picture.width
                if height > 600:
                    width, height = width * 600 / height, 600
                story.append(Image(io.BytesIO(raw), width=width, height=height))
            else:
                story.append(
                    para(
                        (
                            "• "
                            if kind == "bullet"
                            else f"{list_number}. "
                            if kind == "number"
                            else ""
                        )
                        + text,
                        16 if kind == "heading" else 11,
                    )
                )

        def on_page(canvas, doc):
            canvas.setFont("STSong-Light", 9)
            canvas.drawString(56, 810, header[:100])
            canvas.drawString(56, 28, footer[:100])
            canvas.drawRightString(540, 28, str(doc.page))

        SimpleDocTemplate(
            output,
            pagesize=(595.28, 841.89),
            leftMargin=56,
            rightMargin=56,
            topMargin=56,
            bottomMargin=50,
        ).build(story, onFirstPage=on_page, onLaterPages=on_page)
    else:
        raise ValueError("文档格式只能为 docx 或 pdf")
    return output.getvalue()


def spreadsheet(sheets, *, with_values=False):
    import xlsxwriter
    from openpyxl.formula.tokenizer import Tokenizer
    from openpyxl.utils.cell import range_boundaries, get_column_letter

    if not 1 <= len(sheets) <= 20:
        raise ValueError("工作簿需有 1–20 个工作表")
    grid, formats, names = {}, {}, []
    formula_cells = set()
    for sheet in sheets:
        name = str(sheet.get("name", "Sheet"))
        if (
            not name
            or len(name) > 31
            or re.search(r"[\[\]:*?/\\]", name)
            or name.casefold() in [n.casefold() for n in names]
        ):
            raise ValueError("工作表名称无效或重复")
        names.append(name)
        rows = sheet.get("rows", [])
        if len(rows) > 2000 or any(len(row) > 50 for row in rows):
            raise ValueError("每个工作表最多 2000 行、50 列")
        for r, row in enumerate(rows, 1):
            for c, value in enumerate(row, 1):
                if isinstance(value, dict):
                    formula = value.get("formula")
                    if formula:
                        formula_cells.add((name, f"{get_column_letter(c)}{r}"))
                    grid[name, f"{get_column_letter(c)}{r}"] = (
                        "=" + str(formula).lstrip("=")
                        if formula
                        else value.get("value", "")
                    )
                    formats[name, r, c] = value.get("number_format", "")
                else:
                    # Formula intent must be explicit; uploaded strings are never executable.
                    grid[name, f"{get_column_letter(c)}{r}"] = value
    if len(grid) > 20000:
        raise ValueError("工作簿最多 20000 个单元格")
    cached, active = {}, set()

    def numeric(value):
        return (
            float(value)
            if isinstance(value, (int, float)) and not isinstance(value, bool)
            else 0
        )

    def cell(sheet, address, depth=0):
        key = sheet, address.replace("$", "").upper()
        if depth > 100 or key in active:
            raise ValueError("公式循环引用或嵌套过深")
        if key in cached:
            return cached[key]
        value = grid.get(key, 0)
        # Only explicitly marked cells are formulas.
        is_formula = key in formula_cells
        if not is_formula:
            return value
        active.add(key)

        def reference(ref):
            target, area = ref.rsplit("!", 1) if "!" in ref else (sheet, ref)
            target = target.strip("'").replace("''", "'")
            if (
                target not in names
                or "[" in ref
                or not re.fullmatch(
                    r"\$?[A-Za-z]{1,3}\$?\d+(?::\$?[A-Za-z]{1,3}\$?\d+)?", area
                )
            ):
                raise ValueError("公式引用无效，不允许外部链接")
            c1, r1, c2, r2 = range_boundaries(area.replace("$", ""))
            if (r2 - r1 + 1) * (c2 - c1 + 1) > 20000:
                raise ValueError("公式引用范围过大")
            values = [
                cell(target, f"{get_column_letter(c)}{r}", depth + 1)
                for r in range(r1, r2 + 1)
                for c in range(c1, c2 + 1)
            ]
            return values if ":" in area else values[0]

        variables, expression = {}, ""
        for token in Tokenizer(value).items:
            if token.type == "OPERAND" and token.subtype == "RANGE":
                keyname = f"v{len(variables)}"
                variables[keyname] = reference(token.value)
                expression += keyname
            elif token.type == "FUNC":
                expression += token.value.upper()
            elif token.type == "OPERAND" and token.subtype != "NUMBER":
                raise ValueError("公式仅支持数字、引用和基础函数")
            else:
                expression += token.value.replace("^", "**")

        def evaluate(node):
            if isinstance(node, ast.Expression):
                return evaluate(node.body)
            if isinstance(node, ast.Constant) and type(node.value) in {int, float}:
                return node.value
            if isinstance(node, ast.Name) and node.id in variables:
                return variables[node.id]
            if isinstance(node, ast.UnaryOp) and isinstance(
                node.op, (ast.UAdd, ast.USub)
            ):
                return numeric(evaluate(node.operand)) * (
                    -1 if isinstance(node.op, ast.USub) else 1
                )
            if isinstance(node, ast.BinOp):
                a, b = numeric(evaluate(node.left)), numeric(evaluate(node.right))
                if isinstance(node.op, ast.Add):
                    return a + b
                if isinstance(node.op, ast.Sub):
                    return a - b
                if isinstance(node.op, ast.Mult):
                    return a * b
                if isinstance(node.op, ast.Div):
                    return a / b
                if isinstance(node.op, ast.Pow) and abs(b) <= 20 and abs(a) <= 1e9:
                    return a**b
            if (
                isinstance(node, ast.Call)
                and isinstance(node.func, ast.Name)
                and not node.keywords
            ):
                args = [evaluate(arg) for arg in node.args]
                flat = [
                    v
                    for a in args
                    for v in (a if isinstance(a, list) else [a])
                    if type(v) in {int, float}
                ]
                fn = node.func.id
                if fn == "SUM":
                    return sum(flat)
                if fn == "AVERAGE" and flat:
                    return sum(flat) / len(flat)
                if fn == "MIN" and flat:
                    return min(flat)
                if fn == "MAX" and flat:
                    return max(flat)
                if fn == "COUNT":
                    return len(flat)
                if fn == "ABS" and len(flat) == 1:
                    return abs(flat[0])
                if fn == "ROUND" and len(flat) == 2 and abs(flat[1]) <= 20:
                    from decimal import Decimal, ROUND_HALF_UP

                    return float(
                        Decimal(str(flat[0])).quantize(
                            Decimal(10) ** -int(flat[1]), rounding=ROUND_HALF_UP
                        )
                    )
            raise ValueError(
                "暂不支持该公式；支持四则运算、SUM、AVERAGE、MIN、MAX、COUNT、ABS、ROUND"
            )

        result = evaluate(ast.parse(expression, mode="eval"))
        if not isinstance(result, (int, float)) or not math.isfinite(result):
            raise ValueError("公式结果无效")
        active.remove(key)
        cached[key] = result
        return result

    output = io.BytesIO()
    wb = xlsxwriter.Workbook(
        output,
        {"in_memory": True, "strings_to_urls": False, "strings_to_formulas": False},
    )
    header = wb.add_format(
        {
            "bold": True,
            "bg_color": "#E7EDF4",
            "text_wrap": True,
            "valign": "top",
            "font_name": "Microsoft YaHei",
        }
    )
    body = wb.add_format(
        {"text_wrap": True, "valign": "top", "font_name": "Microsoft YaHei"}
    )
    for sheet in sheets:
        name = sheet["name"]
        ws = wb.add_worksheet(name)
        ws.freeze_panes(1, 0)
        ws.set_column(0, 49, 18)
        for r, row in enumerate(sheet["rows"], 1):
            ws.set_row(r - 1, 32 if r == 1 else 30)
            for c, v in enumerate(row, 1):
                address = f"{get_column_letter(c)}{r}"
                value = grid[name, address]
                fmt = header if r == 1 else body
                number_format = formats.get((name, r, c))
                if number_format:
                    fmt = wb.add_format(
                        {
                            "num_format": str(number_format)[:100],
                            "font_name": "Microsoft YaHei",
                        }
                    )
                if isinstance(v, dict) and v.get("formula"):
                    ws.write_formula(r - 1, c - 1, value, fmt, cell(name, address))
                else:
                    if not isinstance(value, (str, int, float, bool, type(None))) or (
                        isinstance(value, float) and not math.isfinite(value)
                    ):
                        raise ValueError("单元格值必须为文字、数字、布尔值或空值")
                    ws.write(r - 1, c - 1, value, fmt)
        if sheet["rows"] and len(sheet["rows"]) > 1:
            ws.autofilter(
                0, 0, len(sheet["rows"]) - 1, max(map(len, sheet["rows"])) - 1
            )
    wb.close()
    if with_values:
        calculations = [{"sheet": name, "cell": cell, "formula": grid[name, cell], "value": cached.get((name, cell))}
                        for name, cell in sorted(formula_cells)]
        return output.getvalue(), calculations
    return output.getvalue()


def text_file(format, content=None, rows=None):
    if (content is None) == (rows is None):
        raise ValueError("content与rows必须且只能提供一项")
    if rows is not None and format != "csv":
        raise ValueError("rows结构化行仅适用于CSV")
    if format == "json":
        value = json.loads(content)
        return json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False).encode(
            "utf-8"
        )
    if format == "csv":
        if rows is None:
            try:
                rows = list(csv.reader(io.StringIO(content, newline=""), strict=True))
            except csv.Error as error:
                raise ValueError(f"CSV引号或转义不合法：{error}，请改用rows结构化行。") from None
        if not rows or len(rows) > 4000 or any(not isinstance(row, list) or len(row) > 50 for row in rows):
            raise ValueError("CSV需有1–4000行，每行最多50列")
        widths = {len(row) for row in rows if row}
        if len(widths) > 1 or not widths or sum(len(row) for row in rows) > 20000:
            raise ValueError("CSV各数据行须与表头列数一致，且最多20000单元格；含逗号/引号的字段请改用rows。")
        rows = [["" if value is None else str(value) for value in row] for row in rows]
        output = io.StringIO(newline="")
        writer = csv.writer(output)
        for row in rows:
            writer.writerow(
                [
                    (
                        "'" + v
                        if v.lstrip().startswith(("=", "+", "-", "@"))
                        and not re.fullmatch(r"[-+]?\d+(?:\.\d+)?", v.strip())
                        else v
                    )
                    for v in row
                ]
            )
        return output.getvalue().encode("utf-8-sig")
    if format not in {"txt", "md"}:
        raise ValueError("文本格式只能为 txt、md、csv、json")
    return content.encode("utf-8")


def pdf_page(raw, page):
    import pymupdf

    with pymupdf.open(stream=raw, filetype="pdf") as pdf:
        if pdf.is_encrypted:
            raise ValueError("加密 PDF 无法预览")
        if not 1 <= page <= len(pdf):
            raise ValueError("页码超出范围")
        current = pdf[page - 1]
        scale = min(1.6, 1800 / max(current.rect.width, current.rect.height))
        pix = current.get_pixmap(matrix=pymupdf.Matrix(scale, scale), alpha=False)
        return {
            "kind": "pdf",
            "page": page,
            "pages": len(pdf),
            "image": "data:image/png;base64,"
            + base64.b64encode(pix.tobytes("png")).decode(),
        }
