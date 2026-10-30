"""Read real-eval evidence and build a self-contained, searchable HTML report."""
from __future__ import annotations

import argparse
from collections import Counter
from datetime import datetime
import html
import json
from pathlib import Path
import re
from zoneinfo import ZoneInfo

ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / ".local-data/evals-20261027"


def load_cases():
    cases = {}
    for name in ["run-1", "run-2"]:
        for file in sorted((BASE / name).glob("case-*.json")):
            row = json.loads(file.read_text(encoding="utf-8"))
            row["evidence_path"] = str(file)
            cases[row["id"]] = row
    return cases


def artifact_content(path, name):
    result = {"path": str(path), "exists": path.exists(), "name": name}
    if not path.exists():
        return result
    result["bytes"] = path.stat().st_size
    extension = Path(name).suffix.lower()
    if extension == ".docx":
        from docx import Document
        doc = Document(path)
        result["text"] = "\n".join(p.text for p in doc.paragraphs)
        result["tables"] = [[list(c.text for c in row.cells) for row in table.rows] for table in doc.tables]
    elif extension == ".xlsx":
        import openpyxl
        # The stored artifact has a .bin suffix; file handles bypass suffix validation.
        with path.open("rb") as stream:
            wb = openpyxl.load_workbook(stream, data_only=False, read_only=True)
            result["sheets"] = {sheet.title: [[c.value for c in row] for row in sheet] for sheet in wb.worksheets}
            result["formula_count"] = sum(1 for sh in wb for row in sh for c in row if c.data_type == "f")
        with path.open("rb") as stream:
            cached = openpyxl.load_workbook(stream, data_only=True, read_only=True)
            result["cached_sheets"] = {sheet.title: [[c.value for c in row] for row in sheet] for sheet in cached.worksheets}
    elif extension == ".pdf":
        from pypdf import PdfReader
        pdf = PdfReader(path)
        result["pages"] = len(pdf.pages)
        result["text"] = "\n".join(page.extract_text() or "" for page in pdf.pages)
    elif extension in {".csv", ".txt", ".json", ".md"}:
        raw = path.read_bytes()
        result["utf8_bom"] = raw.startswith(b"\xef\xbb\xbf")
        result["text"] = raw.decode("utf-8-sig")
    return result


def diff(before, after):
    result = {}
    # Changes in timestamps alone are still useful for detecting rewrites.
    for table in after:
        identity = lambda x: x.get("id") or json.dumps({k: x[k] for k in ["event_id", "occurrence_start", "start"] if k in x}, sort_keys=True) or json.dumps(x, sort_keys=True)
        old = {identity(x): x for x in before.get(table, [])}
        new = {identity(x): x for x in after[table]}
        changed = [x for key, x in new.items() if old.get(key) != x]
        removed = [x for key, x in old.items() if key not in new]
        if changed or removed:
            result[table] = {"created_or_changed": changed, "removed": removed}
    return result


def short(value, limit):
    string = json.dumps(value, ensure_ascii=False, default=str)
    return value if len(string) <= limit else {"excerpt": string[:limit], "truncated": True}


def public_outputs(row):
    """Expose actual teaching/interview text rather than only workflow IDs."""
    outputs = []
    for turn in row["turns"]:
        run = turn["run"]
        kind = run["kind"]
        content = turn["output"]
        if kind == "learning_intake":
            messages = [m for observation in row["observations"]
                        if isinstance(observation.get("value"), dict)
                        for m in observation["value"].get("messages", [])
                        if m.get("role") == "assistant" and m.get("run_id") == run["id"]]
            if messages:
                content = "\n\n".join(m["content"] for m in messages)
        if kind == "lesson":
            lesson_id = run.get("result", {}).get("lesson_id")
            lessons = [observation["value"] for observation in row["observations"]
                       if observation.get("label") == "实际教学内容" and isinstance(observation.get("value"), dict)
                       and observation["value"].get("id") == lesson_id]
            if lessons:
                lesson = lessons[-1]
                content = lesson.get("content", "") + "\n\n练习：\n" + lesson.get("exercise", "")
        if kind == "learning_review":
            reviews = [review for observation in row["observations"]
                       if observation.get("label") == "独立复习记录" and isinstance(observation.get("value"), list)
                       for review in observation["value"] if review.get("id") == run.get("result", {}).get("review_id")]
            if reviews:
                content = reviews[-1].get("exercise", content)
        outputs.append(content)
    return outputs


def prepare():
    folder = BASE / "review-evidence"
    folder.mkdir(exist_ok=True)
    for n, row in load_cases().items():
        runtime = Path(row["evidence_path"]).parent / f"runtime-{row['chapter']:02}"
        run_ids = {turn["run"]["id"] for turn in row["turns"]}
        old = {x["id"] for x in row["before"]["artifacts"]}
        generated = [x for x in row["after"]["artifacts"] if x["id"] not in old or x["run_id"] in run_ids]
        contents = []
        for a in generated:
            try:
                detail = artifact_content(runtime / "attachments/artifacts" / a["storage_name"], a["name"])
                detail.update({"artifact_id": a["id"], "created_at": a["created_at"]})
                contents.append(detail)
            except Exception as error:
                contents.append({"name": a["name"], "read_error": str(error)})
        summary = {k: row[k] for k in ["id", "title", "chapter", "source", "prompts", "expected", "specification", "sent", "notes", "execution_status"]}
        summary["turns"] = []
        for t in row["turns"]:
            summary["turns"].append({"status": t["run"]["status"], "error": t["run"]["error"], "model": t["run"].get("model_config"),
                "output": t["output"], "result": t["run"].get("result"),
                "tool_calls": [e["data"] for e in t["events"] if e["kind"] == "tool"],
                "retrieval": [e["data"] for e in t["events"] if e["kind"] in {"retrieval", "knowledge_retrieval", "learning_continuation"}],
                "event_kinds": list(dict.fromkeys(e["kind"] for e in t["events"])),
                "presentation": t["presentation"]})
        summary["data_changes"] = diff(row["before"], row["after"])
        summary["observations"] = row["observations"]
        summary["generated_artifacts"] = contents
        supplemental = BASE / "supplemental" / f"case-{n:03}.json"
        if supplemental.exists():
            summary["supplemental"] = json.loads(supplemental.read_text(encoding="utf-8"))
        summary["evidence_path"] = row["evidence_path"]
        (folder / f"case-{n:03}.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    print(f"Prepared {len(load_cases())} evidence files", flush=True)


def report(destination):
    folder = Path(destination)
    folder.mkdir(parents=True, exist_ok=True)
    rows = []
    manual = BASE / "manual-verdicts.json"
    overrides = json.loads(manual.read_text(encoding="utf-8")) if manual.exists() else {}
    for n, raw in sorted(load_cases().items()):
        ev = json.loads((BASE / "review-evidence" / f"case-{n:03}.json").read_text(encoding="utf-8"))
        review_path = BASE / "reviews" / f"case-{n:03}.json"
        judgement = json.loads(review_path.read_text(encoding="utf-8")) if review_path.exists() else {"verdict": "待判定", "reason": "尚未完成证据复核", "checks": []}
        judgement.update(overrides.get(str(n), {}))
        # Never convert missing environment or explicitly unexecuted work into a pass.
        if raw["execution_status"] == "not_applicable":
            judgement["verdict"] = "不适用"
            judgement["reason"] = "；".join(raw["notes"]) or judgement.get("reason", "")
        elif raw["execution_status"] == "blocked":
            judgement["verdict"] = "未完成"
        elif raw["execution_status"] == "partial_execution" and judgement["verdict"] == "通过":
            judgement["verdict"] = "部分"
            judgement["reason"] += "；存在尚未执行的原生窗口或界面检查点。"
        questions = [message["content"] + ("\n附件：" + "、".join(message["files"]) if message.get("files") else "")
                     for message in raw["sent"]]
        rows.append({"id": n, "title": raw["title"], "chapter": raw["chapter"],
            "问题": questions or raw["prompts"] or [raw["specification"]],
            "预期答案": raw["expected"], "模型输出": public_outputs(raw),
            "是否通过": judgement["verdict"], "判定依据": judgement.get("reason", ""),
            "逐项检查": judgement.get("checks", []), "执行说明": raw["notes"],
            "实际发送": raw["sent"], "接口与数据观察": raw["observations"],
            "补充验证": ev.get("supplemental", {}), "生成文件检查": ev["generated_artifacts"],
            "工具调用": [{"round": i + 1, "calls": t["tool_calls"]} for i, t in enumerate(ev["turns"])],
            "数据变化": ev["data_changes"], "原始证据": raw["evidence_path"],
            "重试记录": ev.get("supplemental", {}).get("retest", [])})
    assert [r["id"] for r in rows] == list(range(1, 171)), "Report must contain all 170 cases"
    counts = Counter(r["是否通过"] for r in rows)
    meta = {"date": "2026-10-27", "timezone": "Asia/Shanghai", "model": "DeepSeek / deepseek-flash",
            "generated_at": datetime.now(ZoneInfo("Asia/Shanghai")).isoformat(), "total": len(rows), "counts": dict(counts),
            "scope": "复用测试版真实 API 配置，在独立数据库调用当前工作区实际 Agent。通过表示本报告列明检查点有证据；原生弹窗、截图粘贴、实际重启和手机收件未完成的项目不会记为通过。",
            "limits": "Tavily、Embedding、图片生成、MCP、微信未配置；双连接模型快照前置不满足。Embedding 可降级案例按关键词模式执行。判定以有效首轮为主；脚本自身错误修正后重跑的有效记录替换无效记录，功能失败复测单独列出。通过不等同于完整桌面验收；部分包括检查点未满足或尚未验证。所有最终结论均经本次助手复核，模型辅助判读未直接作为最终结果。"}
    (folder / "results.json").write_text(json.dumps({"summary": meta, "cases": rows}, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    def esc(value):
        escaped = html.escape(str(value), quote=True)
        # Preserve original model whitespace in <pre> without trailing whitespace
        # in the tracked HTML source (Markdown hard breaks may contain two spaces).
        return re.sub(r"[ \t]+(?=\n|$)", lambda m: "".join("&#32;" if c == " " else "&#9;" for c in m[0]), escaped)
    json_pre = lambda x: "<pre>" + esc(json.dumps(x, ensure_ascii=False, indent=2, default=str)) + "</pre>"
    cards = []
    for r in rows:
        status = r["是否通过"]
        tone = {"通过": "pass", "失败": "fail", "部分": "partial", "不适用": "na", "未完成": "na"}.get(status, "na")
        questions = "\n\n".join(f"第 {i+1} 条：\n{p}" for i, p in enumerate(r["问题"]))
        output = "\n\n".join(f"第 {i+1} 轮模型输出：\n{p}" for i, p in enumerate(r["模型输出"]))
        if not output:
            output = "未产生模型输出。本例为操作或接口检查，或前置条件不满足。请查看执行说明及接口观察，不能视为模型已回答。"
        checks = "".join(f"<li>{esc(c.get('result',''))} · {esc(c.get('criterion',''))}：{esc(c.get('evidence',''))}</li>" for c in r["逐项检查"])
        notes = "".join(f"<li>{esc(n)}</li>" for n in r["执行说明"])
        retry = ""
        if r["重试记录"]:
            retry = "<details><summary>失败复测记录</summary>" + json_pre(r["重试记录"]) + "</details>"
        cards.append(f'''<article id="case-{r['id']:03}" data-status="{esc(status)}" data-chapter="{r['chapter']}">
<header><h2>{r['id']:03}. {esc(r['title'])}</h2><span class="badge {tone}">{esc(status)}</span></header>
<h3>问题</h3><pre>{esc(questions)}</pre><h3>预期答案</h3><ul>{''.join('<li>'+esc(e)+'</li>' for e in r['预期答案'])}</ul>
<h3>模型输出</h3><pre>{esc(output)}</pre><h3>是否通过：{esc(status)}</h3><p class="reason">{esc(r['判定依据'])}</p>
{('<ul>'+notes+'</ul>') if notes else ''}<details><summary>判定证据</summary>{('<ul>'+checks+'</ul>') if checks else ''}{json_pre({'接口与数据观察':r['接口与数据观察'],'数据变化':r['数据变化'],'补充验证':r['补充验证']})}</details>
<details><summary>实际工具调用与生成文件</summary>{json_pre({'工具调用':r['工具调用'],'生成文件检查':r['生成文件检查']})}</details>{retry}</article>''')
    count_html = "".join(f"<button data-filter='{esc(k)}'>{esc(k)} <strong>{v}</strong></button>" for k, v in counts.items())
    page = '''<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>XiaoXi 170 项真实 API 评测结果</title>
<style>:root{font:16px/1.7 system-ui,"Microsoft YaHei",sans-serif;color:#26352f;background:#f4f5f0}body{max-width:1100px;margin:auto;padding:30px 24px 80px}h1{font-size:29px;margin:0}h2{font-size:21px;margin:0}h3{font-size:16px;margin:20px 0 6px}p{margin:10px 0}a{color:#26664a}article{background:white;border:1px solid #dce4dc;border-radius:14px;margin:22px 0;padding:25px;scroll-margin-top:150px}article header{display:flex;align-items:center;justify-content:space-between;gap:16px}.badge{padding:4px 15px;border-radius:30px;white-space:nowrap;font-weight:650}.pass{background:#e2f3e6;color:#20663c}.fail{background:#fce5df;color:#a33b28}.partial{background:#fff2d3;color:#8c5e09}.na{background:#edf0f4;color:#576373}pre{white-space:pre-wrap;overflow-wrap:anywhere;background:#f6f7f4;padding:16px;border-radius:8px;font:14px/1.8 ui-monospace,"Microsoft YaHei",sans-serif}details{margin-top:16px;border-top:1px solid #e3e7df;padding-top:12px}summary{cursor:pointer;color:#32604a}details pre{max-height:700px;overflow:auto;font-size:12px}.intro{background:#e9eee6;border-radius:12px;padding:18px;margin:20px 0}.controls{position:sticky;top:0;z-index:4;background:#f4f5f0ed;backdrop-filter:blur(12px);padding:12px 0;display:flex;gap:8px;flex-wrap:wrap}button,input,select{font:inherit;border:1px solid #cad4ca;border-radius:8px;background:white;padding:7px 11px}button{cursor:pointer}input{flex:1;min-width:190px}li{margin:5px 0}.reason{font-weight:550}#visible{font-size:13px;color:#667368}.hidden{display:none}@media print{.controls{display:none}article{break-inside:avoid}details{display:none}body{max-width:none;padding:0}pre{font-size:11px}}</style>
<h1>XiaoXi 170 项真实 API 评测结果</h1>'''
    page += f"<p>2026-10-27 · DeepSeek / deepseek-flash · 北京时间 · 共 170 条</p><div class='intro'><p>{esc(meta['scope'])}</p><p>{esc(meta['limits'])}</p><p>所有模型输出按原始响应保留，未用预期答案替换。仅为操作类的案例会展示实际接口数据。失败重试单列，成功的重试不会抹去首轮失败。</p></div>"
    page += "<div class='controls'><button data-filter='全部'>全部 <strong>170</strong></button>" + count_html + "<select id='chapter'><option value=''>全部分册</option>" + "".join(f"<option value='{i}'>第 {i} 册</option>" for i in range(1, 9)) + "</select><input id='search' placeholder='搜索编号、问题或输出…'></div><p id='visible'>显示 170 / 170 条</p>" + "".join(cards)
    page += '''<script>let filter='全部';const articles=[...document.querySelectorAll('article')];function apply(){let q=document.querySelector('#search').value.trim().toLowerCase(),chapter=document.querySelector('#chapter').value,count=0;articles.forEach(a=>{let show=(filter==='全部'||a.dataset.status===filter)&&(!chapter||a.dataset.chapter===chapter)&&(!q||a.textContent.toLowerCase().includes(q)||a.id.includes(q));a.classList.toggle('hidden',!show);if(show)count++});document.querySelector('#visible').textContent=`显示 ${count} / 170 条`;}document.querySelectorAll('[data-filter]').forEach(b=>b.onclick=()=>{filter=b.dataset.filter;apply()});document.querySelector('#search').oninput=apply;document.querySelector('#chapter').onchange=apply;</script></html>'''
    (folder / "评测结果.html").write_text(page, encoding="utf-8")
    print(json.dumps(meta, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--prepare", action="store_true")
    parser.add_argument("--report", default="")
    args = parser.parse_args()
    if args.prepare:
        prepare()
    if args.report:
        report(args.report)
