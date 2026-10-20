"""Create and independently open QA formats in the isolated office workspace."""

import io
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "backend"))
from app import office_formats as fmt
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

root = Path(__file__).resolve().parents[1] / ".local-data/qa-office"
folder = root / "fixtures"
folder.mkdir(parents=True, exist_ok=True)
blocks = [
    {"kind": "heading", "text": "项目进展", "level": 1},
    {
        "kind": "paragraph",
        "text": "项目预算为100元，目前已完成3项任务。下一步整理文档并验收。",
    },
    {
        "kind": "table",
        "rows": [
            ["事项", "预算（元）", "状态"],
            ["资料整理", "40", "完成"],
            ["文档交付", "60", "待验收"],
        ],
    },
    {"kind": "bullet", "text": "原始资料保留，输出文件单独保存。"},
]
for format in ["docx", "pdf"]:
    raw = fmt.document(
        format, "办公文件验收示例", blocks, "XiaoXi 办公处理", "隔离测试资料"
    )
    (folder / f"source.{format}").write_bytes(raw)
    if format == "docx":
        assert len(Document(io.BytesIO(raw)).tables) == 1
    else:
        assert "项目预算为100元" in PdfReader(io.BytesIO(raw)).pages[0].extract_text()
        import base64

        (folder / "pdf-page.png").write_bytes(
            base64.b64decode(fmt.pdf_page(raw, 1)["image"].split(",")[1])
        )
raw = fmt.spreadsheet(
    [
        {
            "name": "明细",
            "rows": [
                ["项目", "金额"],
                ["甲", 40],
                ["乙", 60],
                ["合计", {"formula": "SUM(B2:B3)", "number_format": "0.00"}],
            ],
        }
    ]
)
(folder / "sample.xlsx").write_bytes(raw)
wb = load_workbook(io.BytesIO(raw), data_only=True)
assert wb["明细"]["B4"].value == 100
wb.close()
for format, content in {
    "txt": "中文办公资料\n预算100元",
    "md": "# 项目笔记\n预算100元",
    "csv": "事项,预算\n甲,40\n乙,60",
    "json": '{"budget":100,"completed":3}',
}.items():
    (folder / f"sample.{format}").write_bytes(fmt.text_file(format, content))
for path in folder.iterdir():
    if path.suffix.lstrip(".") in fmt.MIMES:
        assert fmt.parse(path.read_bytes(), path.name)["chunks"]
print(
    json.dumps(
        {"formats": 7, "formula_cache": 100, "folder": str(folder)}, ensure_ascii=True
    )
)
